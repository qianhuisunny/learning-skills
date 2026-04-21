"""Generate a formatted Word doc (.docx) from a content-map.md file.

Usage:
    python3 generate_content_map_docx.py <module_dir>

Reads:  <module_dir>/content-map.md (YAML-ish markdown with O-A-P-P table)
Writes: <module_dir>/content-map.docx
"""
from __future__ import annotations
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


def parse_content_map(md_text: str) -> dict:
    """Extract header fields and table rows from content-map.md."""
    result = {"title": "", "meta": [], "rows": [], "notes": []}

    # Title
    m = re.search(r"^#\s+(.+)$", md_text, re.MULTILINE)
    if m:
        result["title"] = m.group(1).strip()

    # Meta lines (bold key: value)
    for m in re.finditer(r"^\*\*(.+?):\*\*\s*(.+)$", md_text, re.MULTILINE)    :
        result["meta"].append((m.group(1).strip(), m.group(2).strip()))

    # Table rows (skip header + separator)
    table_lines = [l for l in md_text.splitlines()
                   if l.startswith("|") and not re.match(r"^\|[-\s|]+\|$", l)]
    if len(table_lines) > 1:
        for line in table_lines[1:]:  # skip header row
            cells = [c.strip().replace("**", "") for c in line.split("|")[1:-1]]
            if len(cells) >= 4:
                result["rows"].append(cells[:4])

    # Notes section
    notes_match = re.search(r"## Notes\s*\n(.*?)(?=\n##|\Z)", md_text, re.DOTALL)
    if notes_match:
        for line in notes_match.group(1).strip().splitlines():
            line = line.strip().lstrip("- ")
            if line:
                result["notes"].append(line)

    return result


def build_docx(data: dict, output_path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title
    title = doc.add_heading(data["title"], level=1)
    title.runs[0].font.color.rgb = RGBColor(0x1A, 0x23, 0x32)

    # Meta
    for key, val in data["meta"]:
        p = doc.add_paragraph()
        run_key = p.add_run(f"{key}: ")
        run_key.bold = True
        run_key.font.size = Pt(11)
        run_val = p.add_run(val)
        run_val.font.size = Pt(11)
        p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph()  # spacer

    # O-A-P-P table
    headers = ["Objective", "Assessment", "Practice", "Presentation"]
    table = doc.add_table(rows=1 + len(data["rows"]), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row — blue background
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # Set cell shading to brand blue
        from docx.oxml.ns import qn
        shading = cell._element.get_or_add_tcPr()
        shd = shading.makeelement(qn("w:shd"), {
            qn("w:fill"): "1A9AFA",
            qn("w:val"): "clear",
        })
        shading.append(shd)

    # Data rows
    for r_idx, row in enumerate(data["rows"]):
        for c_idx, cell_text in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = cell_text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
            # Alternate row shading
            if r_idx % 2 == 1:
                from docx.oxml.ns import qn
                shading = cell._element.get_or_add_tcPr()
                shd = shading.makeelement(qn("w:shd"), {
                    qn("w:fill"): "F0F4FA",
                    qn("w:val"): "clear",
                })
                shading.append(shd)

    # Set column widths
    widths = [Inches(2.0), Inches(2.0), Inches(1.5), Inches(2.5)]
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

    # Notes section
    if data["notes"]:
        doc.add_paragraph()
        doc.add_heading("Notes", level=2)
        for note in data["notes"]:
            doc.add_paragraph(note, style="List Bullet")

    doc.save(str(output_path))
    return output_path


def main():
    if len(sys.argv) != 2:
        print("Usage: python3 generate_content_map_docx.py <module_dir>", file=sys.stderr)
        sys.exit(2)
    module_dir = Path(sys.argv[1]).expanduser().resolve()
    md_path = module_dir / "content-map.md"
    if not md_path.exists():
        print(f"Error: {md_path} not found", file=sys.stderr)
        sys.exit(1)
    data = parse_content_map(md_path.read_text(encoding="utf-8"))
    out = build_docx(data, module_dir / "content-map.docx")
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
