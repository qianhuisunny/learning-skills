#!/usr/bin/env python3
"""
eLearning Storyboard Template — Dark EY Theme
Reusable python-pptx module for generating storyboard cards.

STYLE RULES:
- Sentence case everywhere (not title case)
- Body content starts at same left margin as title
- Pop-ups are dark panels with blue (#1A9AFA) border, on-theme with the dark slide
- Content inside pop-ups uses white/light text with dash-style bullets

Usage:
    from storyboard_template import StoryboardBuilder
    sb = StoryboardBuilder("Module 2.5 — Developing the procurement plan")
    sb.add_title_slide("Module 2.5", "Developing the procurement plan", "Subtitle")
    sb.add_anchor("2.5-2", "Title", tiles=[...])
    sb.add_popup_text("2.5-2-01", "Anchor title", 1, "Component name", ...)
    sb.save("/path/to/output.pptx")
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ── Dark EY Theme Constants — aligned to theme-dark.css tokens ──
BG_COLOR        = RGBColor(0x1C, 0x23, 0x33)   # --bg-primary
BG_SECONDARY    = RGBColor(0x25, 0x2D, 0x42)   # --bg-secondary
BG_ELEVATED     = RGBColor(0x2E, 0x38, 0x4F)   # --bg-elevated
BRAND_BLUE      = RGBColor(0x1A, 0x9A, 0xFA)   # --brand-blue
BRAND_BLUE_DIM  = RGBColor(0x15, 0x78, 0xC4)   # --brand-blue-dim
WHITE           = RGBColor(0xFF, 0xFF, 0xFF)    # --text-primary
TEXT_SECONDARY  = RGBColor(0xB4, 0xBC, 0xD0)    # --text-secondary
TEXT_MUTED      = RGBColor(0x7A, 0x82, 0x99)    # --text-muted
YELLOW_ACCENT   = RGBColor(0xFF, 0xC8, 0x3D)    # --accent-yellow
SUCCESS_GREEN   = RGBColor(0x2E, 0xCC, 0x71)    # --feedback-correct
ERROR_RED       = RGBColor(0xE7, 0x4C, 0x3C)    # --feedback-incorrect
FOOTER_GRAY     = RGBColor(0x7A, 0x82, 0x99)    # same as --text-muted

# Aliases for backward compatibility
TITLE_BLUE      = BRAND_BLUE
COMPONENT_BLUE  = BRAND_BLUE
DARK_BLUE       = BRAND_BLUE_DIM
BORDER_BLUE     = BRAND_BLUE
PANEL_BG        = BG_SECONDARY
LIGHT_GRAY      = TEXT_SECONDARY
MED_GRAY        = TEXT_MUTED
GREEN_BULLET    = BRAND_BLUE    # bullets are blue dashes now, not green dots
BORDER_LIGHT    = BG_ELEVATED
CHARCOAL        = RGBColor(0x2C, 0x3E, 0x50)
DARK_GRAY       = RGBColor(0x4A, 0x4A, 0x4A)
FAINT_GRAY      = RGBColor(0xF4, 0xF6, 0xF7)

FONT_TITLE = "Calibri"
FONT_BODY  = "Calibri"

SLIDE_WIDTH  = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

LEFT_MARGIN   = Inches(0.7)
CONTENT_WIDTH = Inches(11.9)

# Pop-up dark panel geometry — respects LEFT_MARGIN
POPUP_PANEL_TOP    = Inches(1.15)
POPUP_PANEL_LEFT   = LEFT_MARGIN
POPUP_PANEL_WIDTH  = CONTENT_WIDTH
POPUP_PANEL_HEIGHT = Inches(5.8)
POPUP_PANEL_BOTTOM = POPUP_PANEL_TOP + POPUP_PANEL_HEIGHT

# Inside the dark panel, content margins
POPUP_CONTENT_LEFT = LEFT_MARGIN + Inches(0.3)
POPUP_CONTENT_WIDTH = CONTENT_WIDTH - Inches(0.6)
POPUP_BAR_TOP      = POPUP_PANEL_TOP + Inches(0.1)
POPUP_CONTENT_TOP  = POPUP_PANEL_TOP + Inches(0.8)


class StoryboardBuilder:
    """Builds a storyboard PowerPoint in the dark EY theme."""

    def __init__(self, module_name="eLearning module", date_str="12 March 2026"):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_WIDTH
        self.prs.slide_height = SLIDE_HEIGHT
        self.module_name = module_name
        self.date_str = date_str
        self.slide_count = 0

    # ── Private helpers ──────────────────────────────────────────

    def _new_slide(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = BG_COLOR
        self.slide_count += 1
        return slide

    def _add_footer(self, slide):
        num = self.slide_count
        tb = slide.shapes.add_textbox(Inches(0.3), Inches(7.05), Inches(0.4), Pt(18))
        p = tb.text_frame.paragraphs[0]
        p.text = str(num)
        p.font.size = Pt(9)
        p.font.color.rgb = FOOTER_GRAY
        p.font.name = FONT_BODY
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(1.5), Pt(18))
        p = tb.text_frame.paragraphs[0]
        p.text = self.date_str
        p.font.size = Pt(9)
        p.font.color.rgb = FOOTER_GRAY
        p.font.name = FONT_BODY
        tb = slide.shapes.add_textbox(Inches(3.5), Inches(7.05), Inches(6.5), Pt(18))
        p = tb.text_frame.paragraphs[0]
        p.text = f"Internal use only  |  {self.module_name}"
        p.font.size = Pt(9)
        p.font.color.rgb = FOOTER_GRAY
        p.font.name = FONT_BODY
        p.alignment = PP_ALIGN.CENTER

    def _add_slide_title(self, slide, text, size=Pt(32)):
        tb = slide.shapes.add_textbox(LEFT_MARGIN, Inches(0.4), Inches(12), Pt(50))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = size
        p.font.color.rgb = WHITE
        p.font.bold = True
        p.font.name = FONT_TITLE
        # Blue accent line under title (matches CSS)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                      LEFT_MARGIN, Inches(0.95),
                                      CONTENT_WIDTH, Pt(3))
        line.fill.solid()
        line.fill.fore_color.rgb = BRAND_BLUE
        line.line.fill.background()

    def _add_layout_badge(self, slide, slide_num, layout_id):
        tb = slide.shapes.add_textbox(Inches(10.5), Inches(0.4), Inches(2.5), Pt(25))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = f"{slide_num}  |  {layout_id}"
        p.font.size = Pt(12)
        p.font.color.rgb = LIGHT_GRAY
        p.font.bold = True
        p.font.name = FONT_BODY
        p.alignment = PP_ALIGN.RIGHT

    def _add_speaker_notes(self, slide, notes):
        slide.notes_slide.notes_text_frame.text = notes

    def _add_section_label(self, slide, text, left, top, width=Inches(4),
                           color=YELLOW_ACCENT):
        tb = slide.shapes.add_textbox(left, top, width, Pt(18))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(10)
        p.font.color.rgb = color
        p.font.bold = True
        p.font.name = FONT_BODY

    def _add_content_box(self, slide, left, top, width, height,
                         fill_color=PANEL_BG, border_color=BORDER_BLUE):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = fill_color
        box.line.color.rgb = border_color
        box.line.width = Pt(1)
        return box

    def _add_text(self, slide, text, left, top, width, height,
                  size=Pt(16), color=WHITE, bold=False, align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = size
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = FONT_BODY
        p.alignment = align
        p.line_spacing = Pt(size.pt * 1.5)
        return tf

    def _add_bullets(self, slide, items, left, top, width, height,
                     size=Pt(15), bullet_color=BRAND_BLUE, text_color=WHITE):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r1 = p.add_run()
            r1.text = "\u2013  "
            r1.font.size = Pt(14)
            r1.font.color.rgb = bullet_color
            r1.font.name = FONT_BODY
            r2 = p.add_run()
            r2.text = item
            r2.font.size = size
            r2.font.color.rgb = text_color
            r2.font.name = FONT_BODY
            p.space_after = Pt(10)
            p.line_spacing = Pt(24)
        return tf

    def _add_dash_bullets(self, slide, items, left, top, width, height,
                          size=Pt(15), dash_color=COMPONENT_BLUE,
                          text_color=WHITE):
        """Dash-style bullets — used inside dark pop-up panels."""
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r1 = p.add_run()
            r1.text = "\u2013  "
            r1.font.size = Pt(14)
            r1.font.color.rgb = dash_color
            r1.font.name = FONT_BODY
            r2 = p.add_run()
            r2.text = item
            r2.font.size = size
            r2.font.color.rgb = text_color
            r2.font.name = FONT_BODY
            p.space_after = Pt(10)
            p.line_spacing = Pt(24)
        return tf

    def _add_media_placeholder(self, slide, label, left, top, width, height,
                               on_white=False):
        fill = WHITE if on_white else PANEL_BG
        border = BORDER_LIGHT if on_white else BORDER_BLUE
        self._add_content_box(slide, left, top, width, height,
                              fill_color=fill, border_color=border)
        txt_color = MED_GRAY if on_white else LIGHT_GRAY
        self._add_text(slide, label, left + Inches(0.2), top + Inches(0.2),
                       width - Inches(0.4), height - Inches(0.4),
                       size=Pt(13), color=txt_color, align=PP_ALIGN.CENTER)

    def _format_notes(self, vo="", visual="", interaction=""):
        parts = []
        if vo:
            parts.append(f"VOICEOVER:\n{vo}")
        if visual:
            parts.append(f"VISUAL/GRAPHIC DESCRIPTION:\n{visual}")
        if interaction:
            parts.append(f"INTERACTION & DEV NOTES:\n{interaction}")
        return "\n\n".join(parts)

    def _format_feedback_notes(self, vo="", correct="", incorrect="", dev=""):
        parts = []
        if vo:
            parts.append(f"VOICEOVER:\n{vo}")
        if correct:
            parts.append(f"FEEDBACK — CORRECT:\n{correct}")
        if incorrect:
            parts.append(f"FEEDBACK — INCORRECT:\n{incorrect}")
        if dev:
            parts.append(f"DEV NOTES:\n{dev}")
        return "\n\n".join(parts)

    # ── Pop-up panel helpers ─────────────────────────────────────

    def _add_popup_dark_panel(self, slide):
        """Draw the dark panel with blue border that IS the pop-up."""
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            POPUP_PANEL_LEFT, POPUP_PANEL_TOP,
            POPUP_PANEL_WIDTH, POPUP_PANEL_HEIGHT)
        panel.fill.solid()
        panel.fill.fore_color.rgb = PANEL_BG
        panel.line.color.rgb = BRAND_BLUE
        panel.line.width = Pt(2)
        return panel

    def _add_popup_header(self, slide, comp_num, comp_title):
        """Popup header matching CSS: blue num + title + X close + separator line."""
        hdr_y = POPUP_BAR_TOP
        hdr_h = Inches(0.5)
        hdr_x = POPUP_PANEL_LEFT + Inches(0.3)
        hdr_w = POPUP_PANEL_WIDTH - Inches(0.6)
        # Number: blue text, no background
        self._add_text(slide, f"{comp_num:02d}", hdr_x, hdr_y,
                       Inches(0.6), hdr_h, size=Pt(18),
                       color=BRAND_BLUE, bold=True)
        # Title text
        self._add_text(slide, comp_title,
                       hdr_x + Inches(0.6), hdr_y,
                       hdr_w - Inches(1.2), hdr_h,
                       size=Pt(16), color=WHITE, bold=True)
        # X close button: text only, no background
        self._add_text(slide, "\u2715",
                       hdr_x + hdr_w - Inches(0.4), hdr_y,
                       Inches(0.4), hdr_h,
                       size=Pt(18), color=TEXT_SECONDARY, bold=False,
                       align=PP_ALIGN.CENTER)
        # Separator line below header (CSS: border-bottom: 1px solid --bg-elevated)
        sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     hdr_x, hdr_y + hdr_h + Inches(0.05),
                                     hdr_w, Pt(1))
        sep.fill.solid()
        sep.fill.fore_color.rgb = BG_ELEVATED
        sep.line.fill.background()

    # ── Public: Title Slide ──────────────────────────────────────

    def add_title_slide(self, module_num, title, subtitle=""):
        slide = self._new_slide()
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     Inches(0), Inches(0), SLIDE_WIDTH, Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = YELLOW_ACCENT
        bar.line.fill.background()
        self._add_text(slide, module_num.upper(), LEFT_MARGIN, Inches(2.2),
                       Inches(10), Pt(30), size=Pt(18), color=LIGHT_GRAY, bold=True)
        self._add_text(slide, title, LEFT_MARGIN, Inches(2.8),
                       Inches(10), Inches(1.2), size=Pt(44), color=WHITE, bold=True)
        if subtitle:
            self._add_text(slide, subtitle, LEFT_MARGIN, Inches(4.2),
                           Inches(10), Pt(30), size=Pt(18), color=LIGHT_GRAY)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L01 — 1C Text ────────────────────────────

    def add_L01(self, slide_num, title, body, bullets=None, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L01")
        self._add_text(slide, body, LEFT_MARGIN, Inches(1.4),
                       CONTENT_WIDTH, Inches(2), size=Pt(18))
        if bullets:
            self._add_bullets(slide, bullets, LEFT_MARGIN, Inches(3.0),
                              CONTENT_WIDTH, Inches(3.5))
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L02 — 1C Hero + Body ─────────────────────

    def add_L02(self, slide_num, title, hero_label, body, cta=None, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L02")
        self._add_media_placeholder(slide, hero_label, LEFT_MARGIN, Inches(1.3),
                                    CONTENT_WIDTH, Inches(3))
        self._add_text(slide, body, LEFT_MARGIN, Inches(4.6),
                       CONTENT_WIDTH, Inches(1.5), size=Pt(18))
        if cta:
            btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                         Inches(10.5), Inches(6.2), Inches(2), Inches(0.5))
            btn.fill.solid()
            btn.fill.fore_color.rgb = COMPONENT_BLUE
            btn.line.fill.background()
            self._add_text(slide, cta, Inches(10.5), Inches(6.2),
                           Inches(2), Inches(0.5), size=Pt(14),
                           color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L03 — 2C 60/40 (Media Left) ──────────────

    def add_L03(self, slide_num, title, media_label, heading, body,
                bullets=None, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L03")
        self._add_media_placeholder(slide, media_label, LEFT_MARGIN, Inches(1.4),
                                    Inches(7), Inches(4.8))
        rx = Inches(8.0)
        rw = Inches(4.8)
        self._add_text(slide, heading, rx, Inches(1.4), rw, Pt(30),
                       size=Pt(20), color=WHITE, bold=True)
        self._add_text(slide, body, rx, Inches(2.0), rw, Inches(2), size=Pt(16))
        if bullets:
            self._add_bullets(slide, bullets, rx, Inches(3.5), rw, Inches(3))
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L04 — 2C 40/60 (Media Right) ─────────────

    def add_L04(self, slide_num, title, heading, body, media_label,
                bullets=None, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L04")
        lw = Inches(4.8)
        self._add_text(slide, heading, LEFT_MARGIN, Inches(1.4), lw, Pt(30),
                       size=Pt(20), color=WHITE, bold=True)
        self._add_text(slide, body, LEFT_MARGIN, Inches(2.0), lw, Inches(2), size=Pt(16))
        if bullets:
            self._add_bullets(slide, bullets, LEFT_MARGIN, Inches(3.5), lw, Inches(3))
        self._add_media_placeholder(slide, media_label, Inches(5.8), Inches(1.4),
                                    Inches(7), Inches(4.8))
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L05 — 3-up Cards ──────────────────────────

    def add_L05(self, slide_num, title, cards, notes=""):
        """cards: list of (icon_label, heading, body) tuples, max 3."""
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L05")
        card_w = Inches(3.7)
        gap = Inches(0.4)
        for i, (icon, heading, body) in enumerate(cards[:3]):
            x = LEFT_MARGIN + i * (card_w + gap)
            y = Inches(1.8)
            self._add_content_box(slide, x, y, card_w, Inches(4))
            self._add_text(slide, icon, x, y + Inches(0.3), card_w, Pt(20),
                           size=Pt(12), color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
            self._add_text(slide, heading, x + Inches(0.2), y + Inches(1.0),
                           card_w - Inches(0.4), Pt(30),
                           size=Pt(18), color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            self._add_text(slide, body, x + Inches(0.2), y + Inches(1.8),
                           card_w - Inches(0.4), Inches(2),
                           size=Pt(14), color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L06 — Compare 2-up ────────────────────────

    def add_L06(self, slide_num, title, col_a_header, col_a_items,
                col_b_header, col_b_items, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L06")
        col_w = Inches(5.8)
        ax = LEFT_MARGIN
        self._add_content_box(slide, ax, Inches(1.6), col_w, Inches(0.5))
        self._add_text(slide, col_a_header, ax + Inches(0.2), Inches(1.6),
                       col_w, Inches(0.5), size=Pt(16), color=WHITE, bold=True)
        self._add_bullets(slide, col_a_items, ax + Inches(0.2), Inches(2.3),
                          col_w - Inches(0.4), Inches(4))
        bx = Inches(6.8)
        self._add_content_box(slide, bx, Inches(1.6), col_w, Inches(0.5))
        self._add_text(slide, col_b_header, bx + Inches(0.2), Inches(1.6),
                       col_w, Inches(0.5), size=Pt(16), color=WHITE, bold=True)
        self._add_bullets(slide, col_b_items, bx + Inches(0.2), Inches(2.3),
                          col_w - Inches(0.4), Inches(4))
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L07 — Stepper (1→N) ──────────────────────

    def add_L07(self, slide_num, title, steps, notes=""):
        """steps: list of (label, caption) tuples."""
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L07")
        n = len(steps)
        gap = Inches(0.25 if n > 5 else 0.8)
        step_w = Emu((CONTENT_WIDTH - (n - 1) * gap) // n)
        circle_size = Inches(0.45 if n > 5 else 0.5)
        start_x = LEFT_MARGIN
        y_circle = Inches(2.5)
        for i, (label, caption) in enumerate(steps):
            x = start_x + i * (step_w + gap)
            circle_x = x + Emu((step_w - circle_size) // 2)
            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                            circle_x, y_circle,
                                            circle_size, circle_size)
            circle.fill.solid()
            circle.fill.fore_color.rgb = COMPONENT_BLUE
            circle.line.fill.background()
            self._add_text(slide, str(i + 1), circle_x, y_circle,
                           circle_size, circle_size, size=Pt(14 if n > 5 else 16),
                           color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            if i < n - 1:
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                               x + step_w, y_circle + Inches(0.1),
                                               gap, Inches(0.3))
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = LIGHT_GRAY
                arrow.line.fill.background()
            self._add_text(slide, label, x, y_circle + Inches(0.7),
                           step_w, Pt(34), size=Pt(13 if n > 5 else 16),
                           color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            self._add_text(slide, caption, x, y_circle + Inches(1.2),
                           step_w, Inches(1.7), size=Pt(10 if n > 5 else 13),
                           color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L08 — Timeline ────────────────────────────

    def add_L08(self, slide_num, title, nodes, notes=""):
        """nodes: list of (date_label, content) tuples."""
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L08")
        line_x = Inches(3.0)
        start_y = Inches(1.8)
        gap = Inches(1.1)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                      line_x, start_y,
                                      Pt(2), Inches(len(nodes) * 1.1))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_BLUE
        line.line.fill.background()
        for i, (date_label, content) in enumerate(nodes):
            y = start_y + i * gap
            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                            line_x - Inches(0.1), y,
                                            Inches(0.25), Inches(0.25))
            circle.fill.solid()
            circle.fill.fore_color.rgb = COMPONENT_BLUE
            circle.line.fill.background()
            self._add_text(slide, date_label, LEFT_MARGIN, y,
                           Inches(2), Pt(20), size=Pt(14),
                           color=YELLOW_ACCENT, bold=True, align=PP_ALIGN.RIGHT)
            self._add_text(slide, content, Inches(3.5), y,
                           Inches(8), Inches(0.8), size=Pt(15))
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L09 — MCQ (Default State) ────────────────

    def add_L09_default(self, slide_num, title, stem, options, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L09")
        self._add_text(slide, stem, LEFT_MARGIN, Inches(1.5),
                       CONTENT_WIDTH, Inches(1), size=Pt(18), bold=True)
        y = Inches(2.8)
        for i, opt in enumerate(options):
            letter = chr(65 + i)
            tb = slide.shapes.add_textbox(LEFT_MARGIN + Inches(0.3), y, Inches(10), Pt(30))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r1 = p.add_run()
            r1.text = f"\u25CB  {letter}.  "
            r1.font.size = Pt(16)
            r1.font.color.rgb = COMPONENT_BLUE
            r1.font.name = FONT_BODY
            r2 = p.add_run()
            r2.text = opt
            r2.font.size = Pt(16)
            r2.font.color.rgb = WHITE
            r2.font.name = FONT_BODY
            y += Inches(0.6)
        btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     Inches(10.5), Inches(6.0), Inches(2), Inches(0.5))
        btn.fill.solid()
        btn.fill.fore_color.rgb = COMPONENT_BLUE
        btn.line.fill.background()
        self._add_text(slide, "Submit", Inches(10.5), Inches(6.0),
                       Inches(2), Inches(0.5), size=Pt(14),
                       color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L09 — MCQ (Feedback State) ───────────────

    def add_L09_feedback(self, slide_num, title, stem, options, correct_idx,
                         feedback_correct, feedback_incorrect, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L09 — Feedback")
        lw = Inches(8.2)
        self._add_text(slide, stem, LEFT_MARGIN, Inches(1.5),
                       lw, Inches(1), size=Pt(18), bold=True)
        y = Inches(2.8)
        for i, opt in enumerate(options):
            letter = chr(65 + i)
            is_correct = (i == correct_idx)
            tb = slide.shapes.add_textbox(LEFT_MARGIN + Inches(0.3), y, Inches(7), Pt(30))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r1 = p.add_run()
            r1.text = f"\u25C9  {letter}.  " if is_correct else f"\u25CB  {letter}.  "
            r1.font.size = Pt(16)
            r1.font.color.rgb = SUCCESS_GREEN if is_correct else LIGHT_GRAY
            r1.font.name = FONT_BODY
            r2 = p.add_run()
            r2.text = opt
            r2.font.size = Pt(16)
            r2.font.color.rgb = WHITE
            r2.font.name = FONT_BODY
            y += Inches(0.6)
        panel_x = Inches(8.8)
        panel_w = Inches(4.2)
        self._add_content_box(slide, panel_x, Inches(1.2), panel_w, Inches(2.3))
        self._add_section_label(slide, "\u2713 CORRECT", panel_x + Inches(0.2), Inches(1.35))
        self._add_text(slide, feedback_correct, panel_x + Inches(0.2), Inches(1.7),
                       panel_w - Inches(0.4), Inches(1.5), size=Pt(14))
        self._add_content_box(slide, panel_x, Inches(3.7), panel_w, Inches(2.3))
        self._add_section_label(slide, "\u2717 INCORRECT", panel_x + Inches(0.2), Inches(3.85))
        self._add_text(slide, feedback_incorrect, panel_x + Inches(0.2), Inches(4.2),
                       panel_w - Inches(0.4), Inches(1.5), size=Pt(14))
        btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     panel_x + Inches(1.0), Inches(6.2),
                                     Inches(2), Inches(0.5))
        btn.fill.solid()
        btn.fill.fore_color.rgb = COMPONENT_BLUE
        btn.line.fill.background()
        self._add_text(slide, "Continue", panel_x + Inches(1.0), Inches(6.2),
                       Inches(2), Inches(0.5), size=Pt(14),
                       color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L10 — Scenario Choice (Default) ──────────

    def add_L10_default(self, slide_num, title, context, choices, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L10")
        self._add_content_box(slide, LEFT_MARGIN, Inches(1.3), CONTENT_WIDTH, Inches(2.2))
        self._add_section_label(slide, "SCENARIO", LEFT_MARGIN + Inches(0.2), Inches(1.4))
        self._add_text(slide, context, LEFT_MARGIN + Inches(0.2), Inches(1.8),
                       CONTENT_WIDTH - Inches(0.4), Inches(1.5), size=Pt(16))
        n = len(choices)
        btn_w = Inches(3.5)
        gap = Inches(0.5)
        total = n * btn_w + (n - 1) * gap
        start_x = LEFT_MARGIN
        for i, choice in enumerate(choices):
            x = start_x + i * (btn_w + gap)
            btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                         x, Inches(4.2), btn_w, Inches(1.2))
            btn.fill.solid()
            btn.fill.fore_color.rgb = PANEL_BG
            btn.line.color.rgb = BORDER_BLUE
            btn.line.width = Pt(1)
            self._add_text(slide, choice, x + Inches(0.2), Inches(4.3),
                           btn_w - Inches(0.4), Inches(1),
                           size=Pt(15), align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L10 — Scenario Choice (Feedback) ─────────

    def add_L10_feedback(self, slide_num, title, context, chosen,
                         consequence, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L10 — Feedback")
        lw = Inches(8.2)
        self._add_content_box(slide, LEFT_MARGIN, Inches(1.3), lw, Inches(2.2))
        self._add_section_label(slide, "SCENARIO", LEFT_MARGIN + Inches(0.2), Inches(1.4))
        self._add_text(slide, context, LEFT_MARGIN + Inches(0.2), Inches(1.8),
                       lw - Inches(0.4), Inches(1.5), size=Pt(16))
        self._add_section_label(slide, "SELECTED", LEFT_MARGIN + Inches(0.2), Inches(3.8))
        self._add_text(slide, chosen, LEFT_MARGIN + Inches(0.2), Inches(4.2),
                       lw - Inches(0.4), Inches(1), size=Pt(16), bold=True)
        panel_x = Inches(8.8)
        panel_w = Inches(4.2)
        self._add_content_box(slide, panel_x, Inches(1.2), panel_w, Inches(5))
        self._add_section_label(slide, "CONSEQUENCE", panel_x + Inches(0.2), Inches(1.35))
        self._add_text(slide, consequence, panel_x + Inches(0.2), Inches(1.8),
                       panel_w - Inches(0.4), Inches(4), size=Pt(15))
        btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     panel_x + Inches(1.0), Inches(6.2),
                                     Inches(2), Inches(0.5))
        btn.fill.solid()
        btn.fill.fore_color.rgb = COMPONENT_BLUE
        btn.line.fill.background()
        self._add_text(slide, "Continue", panel_x + Inches(1.0), Inches(6.2),
                       Inches(2), Inches(0.5), size=Pt(14),
                       color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L11 — Drag/Match (Default) ───────────────

    def add_L11_default(self, slide_num, title, instruction, items, targets, notes=""):
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L11")
        self._add_content_box(slide, LEFT_MARGIN, Inches(1.3), CONTENT_WIDTH, Inches(0.5))
        self._add_text(slide, instruction, LEFT_MARGIN + Inches(0.2), Inches(1.3),
                       CONTENT_WIDTH - Inches(0.4), Inches(0.5), size=Pt(14), bold=True)
        self._add_section_label(slide, "DRAG ITEMS", LEFT_MARGIN, Inches(2.1))
        y = Inches(2.5)
        for item in items:
            self._add_content_box(slide, LEFT_MARGIN, y, Inches(4.5), Inches(0.55))
            self._add_text(slide, item, LEFT_MARGIN + Inches(0.2), y,
                           Inches(4.1), Inches(0.55), size=Pt(14))
            y += Inches(0.7)
        self._add_section_label(slide, "DROP TARGETS", Inches(5.8), Inches(2.1))
        y = Inches(2.5)
        for target in targets:
            box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         Inches(5.8), y, Inches(7), Inches(0.55))
            box.fill.solid()
            box.fill.fore_color.rgb = BG_COLOR
            box.line.color.rgb = LIGHT_GRAY
            box.line.width = Pt(1)
            self._add_text(slide, target, Inches(6.0), y,
                           Inches(6.6), Inches(0.55), size=Pt(14), color=LIGHT_GRAY)
            y += Inches(0.7)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Anchor Slide (L12 Menu Grid) ─────────────────────

    def _build_anchor_bars(self, slide, slide_num, title, subtitle, tiles,
                           visited=None):
        self._add_slide_title(slide, title)
        badge_label = "L12" if not visited else "L12 — Visited"
        self._add_layout_badge(slide, slide_num, badge_label)
        if subtitle:
            tb = slide.shapes.add_textbox(LEFT_MARGIN, Inches(1.0), Inches(10), Pt(22))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = subtitle
            p.font.size = Pt(13)
            p.font.color.rgb = LIGHT_GRAY
            p.font.italic = True
            p.font.name = FONT_BODY
        visited = visited or set()
        n = len(tiles)
        cols = min(n, 4)
        tile_w = Inches(2.8)
        tile_h = Inches(1.0)
        gap_x = Inches(0.35)
        gap_y = Inches(0.3)
        # Align grid to LEFT_MARGIN, same as title
        start_x = LEFT_MARGIN
        start_y = Inches(1.6)
        for i, (icon, label) in enumerate(tiles):
            col = i % cols
            row = i // cols
            x = start_x + col * (tile_w + gap_x)
            y = start_y + row * (tile_h + gap_y)
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, tile_w, tile_h)
            bar.fill.solid()
            bar.fill.fore_color.rgb = COMPONENT_BLUE
            bar.line.fill.background()
            badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                           x, y, Inches(0.5), tile_h)
            badge.fill.solid()
            badge.fill.fore_color.rgb = DARK_BLUE
            badge.line.fill.background()
            self._add_text(slide, f"{i+1:02d}", x, y, Inches(0.5), tile_h,
                           size=Pt(14), color=WHITE, bold=True, align=PP_ALIGN.CENTER)
            self._add_text(slide, label, x + Inches(0.6), y,
                           tile_w - Inches(0.7), tile_h,
                           size=Pt(12), color=WHITE, bold=True)
            if i in visited:
                self._add_text(slide, "\u2713", x + tile_w - Inches(0.4), y,
                               Inches(0.35), tile_h,
                               size=Pt(18), color=SUCCESS_GREEN, bold=True,
                               align=PP_ALIGN.CENTER)

    def add_anchor(self, slide_num, title, tiles, subtitle="", notes=""):
        slide = self._new_slide()
        self._build_anchor_bars(slide, slide_num, title, subtitle, tiles)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    def add_anchor_visited(self, slide_num, title, tiles, visited=None,
                           subtitle="", notes=""):
        slide = self._new_slide()
        if visited is None:
            visited = set(range(len(tiles)))
        self._build_anchor_bars(slide, slide_num, title, subtitle, tiles,
                                visited=visited)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Timeline Anchor (L12 horizontal) ────────────────

    def add_timeline_anchor(self, slide_num, title, steps, subtitle="",
                             visited=None, images=None, notes=""):
        """L12 as horizontal timeline — image placeholders above numbered
        circles on a connecting line, step labels below."""
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        badge_text = "L12 — Timeline" + (" visited" if visited else "")
        self._add_layout_badge(slide, slide_num, badge_text)
        if subtitle:
            tb = slide.shapes.add_textbox(LEFT_MARGIN, Inches(1.05),
                                          Inches(10), Pt(22))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = subtitle
            p.font.size = Pt(13)
            p.font.color.rgb = LIGHT_GRAY
            p.font.italic = True
            p.font.name = FONT_BODY
        visited = visited or set()
        n = len(steps)
        cd = Inches(0.42)
        cr = cd // 2
        img_h = Inches(1.1)
        img_w = Inches(1.3)
        group_h = img_h + Inches(0.2) + cd + Inches(0.1) + Inches(0.7)
        avail_top = Inches(1.45)
        avail_bot = Inches(6.9)
        img_y = avail_top + (avail_bot - avail_top - group_h) // 2
        y_top = img_y + img_h + Inches(0.2)
        y_mid = y_top + cr
        total_span = Inches(10.2)
        gap = total_span // (n - 1) if n > 1 else 0
        x0 = (SLIDE_WIDTH - total_span) // 2 - cr
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            x0 + cr, y_mid - Pt(1), gap * (n - 1), Pt(2))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(0x15, 0x5A, 0x8A)
        line.line.fill.background()
        for i, (num_label, title_text) in enumerate(steps):
            cx = x0 + i * gap
            is_v = i in visited
            img_label = images[i] if images and i < len(images) else f"{{{{IMG_{i+1}}}}}"
            img_x = cx + cr - img_w // 2
            self._add_media_placeholder(slide, img_label,
                                        img_x, img_y, img_w, img_h)
            circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx, y_top, cd, cd)
            if is_v:
                circ.fill.solid()
                circ.fill.fore_color.rgb = SUCCESS_GREEN
                circ.line.color.rgb = SUCCESS_GREEN
            else:
                circ.fill.solid()
                circ.fill.fore_color.rgb = BG_SECONDARY
                circ.line.color.rgb = BRAND_BLUE
            circ.line.width = Pt(2)
            nc = WHITE if is_v else BRAND_BLUE
            tb = slide.shapes.add_textbox(cx, y_top, cd, cd)
            tf = tb.text_frame
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.text = num_label
            p.font.size = Pt(12)
            p.font.color.rgb = nc
            p.font.bold = True
            p.font.name = FONT_TITLE
            p.alignment = PP_ALIGN.CENTER
            lw = gap if gap else Inches(1.5)
            lx = cx + cr - lw // 2
            self._add_text(slide, title_text,
                           lx, y_top + cd + Inches(0.1), lw, Inches(1.2),
                           size=Pt(11), color=TEXT_SECONDARY,
                           align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Pop-up Slides (dark panel + blue border) ─────────

    def add_popup_text(self, slide_num, anchor_title, comp_num, comp_title,
                       body="", bullets=None, key_message="", ey_role="",
                       notes=""):
        """Pop-up: dark panel with text + dash bullets + EY Role + key message."""
        slide = self._new_slide()
        self._add_slide_title(slide, anchor_title)
        self._add_layout_badge(slide, slide_num, "Pop-up on anchor")
        self._add_popup_dark_panel(slide)
        self._add_popup_header(slide, comp_num, comp_title)
        ct = POPUP_CONTENT_TOP
        if body:
            self._add_text(slide, body, POPUP_CONTENT_LEFT, ct,
                           POPUP_CONTENT_WIDTH, Inches(1.2),
                           size=Pt(16), color=WHITE)
            ct += Inches(1.0)
        if bullets:
            self._add_dash_bullets(slide, bullets, POPUP_CONTENT_LEFT, ct,
                                    POPUP_CONTENT_WIDTH, Inches(3.5),
                                    size=Pt(15), text_color=WHITE)
        if ey_role:
            ey_y = POPUP_PANEL_BOTTOM - Inches(1.7)
            self._add_section_label(slide, "EY ROLE",
                                    POPUP_CONTENT_LEFT, ey_y,
                                    color=BRAND_BLUE)
            self._add_text(slide, ey_role,
                           POPUP_CONTENT_LEFT + Inches(0.2), ey_y + Inches(0.25),
                           POPUP_CONTENT_WIDTH - Inches(0.2), Inches(0.6),
                           size=Pt(13), color=TEXT_SECONDARY)
        if key_message:
            km_y = POPUP_PANEL_BOTTOM - Inches(0.85)
            km_bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                POPUP_CONTENT_LEFT, km_y,
                POPUP_CONTENT_WIDTH, Inches(0.55))
            km_bar.fill.solid()
            km_bar.fill.fore_color.rgb = RGBColor(0x33, 0x2D, 0x1A)
            km_bar.line.fill.background()
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                POPUP_CONTENT_LEFT, km_y, Pt(4), Inches(0.55))
            accent.fill.solid()
            accent.fill.fore_color.rgb = YELLOW_ACCENT
            accent.line.fill.background()
            self._add_text(slide, key_message,
                           POPUP_CONTENT_LEFT + Inches(0.2), km_y,
                           POPUP_CONTENT_WIDTH - Inches(0.4), Inches(0.55),
                           size=Pt(13), color=YELLOW_ACCENT, bold=True)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    def add_popup_text_media(self, slide_num, anchor_title, comp_num, comp_title,
                             bullets=None, media_label="", body="", notes=""):
        """Pop-up: dark panel with text/bullets left + media placeholder right."""
        slide = self._new_slide()
        self._add_slide_title(slide, anchor_title)
        self._add_layout_badge(slide, slide_num, "Pop-up on anchor")
        self._add_popup_dark_panel(slide)
        self._add_popup_header(slide, comp_num, comp_title)
        ct = POPUP_CONTENT_TOP
        lw = Inches(6)
        if body:
            self._add_text(slide, body, POPUP_CONTENT_LEFT, ct,
                           lw, Inches(1), size=Pt(15), color=WHITE)
            ct += Inches(0.8)
        if bullets:
            self._add_dash_bullets(slide, bullets, POPUP_CONTENT_LEFT, ct,
                                    lw, Inches(3.5), size=Pt(14),
                                    text_color=WHITE)
        if media_label:
            self._add_media_placeholder(slide, media_label,
                                        Inches(7.5), POPUP_CONTENT_TOP,
                                        Inches(5), Inches(4))
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    def add_popup_stepper(self, slide_num, popup_title, comp_num, comp_title,
                          bullets=None, steps=None, extra_sections=None,
                          notes=""):
        """Pop-up: dark panel with bullets + stepper/process flow."""
        slide = self._new_slide()
        self._add_slide_title(slide, popup_title)
        self._add_layout_badge(slide, slide_num, "Pop-up on anchor")
        self._add_popup_dark_panel(slide)
        self._add_popup_header(slide, comp_num, comp_title)
        ct = POPUP_CONTENT_TOP
        if bullets:
            self._add_dash_bullets(slide, bullets, POPUP_CONTENT_LEFT, ct,
                                    Inches(6), Inches(3), size=Pt(13),
                                    text_color=WHITE)
        if steps:
            sx = Inches(7.0)
            sy = POPUP_CONTENT_TOP
            step_w = Inches(1.4)
            for i, (label, caption) in enumerate(steps):
                x = sx + i * (step_w + Inches(0.3))
                box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                             x, sy, step_w, Inches(0.8))
                box.fill.solid()
                box.fill.fore_color.rgb = COMPONENT_BLUE
                box.line.fill.background()
                self._add_text(slide, label, x, sy, step_w, Inches(0.8),
                               size=Pt(12), color=WHITE, bold=True,
                               align=PP_ALIGN.CENTER)
                if i < len(steps) - 1:
                    arrow = slide.shapes.add_shape(
                        MSO_SHAPE.RIGHT_ARROW,
                        x + step_w + Inches(0.02), sy + Inches(0.25),
                        Inches(0.26), Inches(0.3))
                    arrow.fill.solid()
                    arrow.fill.fore_color.rgb = MED_GRAY
                    arrow.line.fill.background()
            for i, (label, caption) in enumerate(steps):
                x = sx + i * (step_w + Inches(0.3))
                self._add_text(slide, caption, x, sy + Inches(0.9),
                               step_w, Inches(0.6), size=Pt(10),
                               color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
        if extra_sections:
            ey = Inches(4.8)
            for sec_title, sec_items in extra_sections:
                self._add_section_label(slide, sec_title,
                                        POPUP_CONTENT_LEFT, ey,
                                        color=COMPONENT_BLUE)
                ey += Inches(0.3)
                for item in sec_items:
                    self._add_text(slide, item, POPUP_CONTENT_LEFT + Inches(0.3), ey,
                                   Inches(11), Pt(18), size=Pt(12), color=LIGHT_GRAY)
                    ey += Inches(0.3)
                ey += Inches(0.15)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    def add_popup_compare(self, slide_num, popup_title, comp_num, comp_title,
                          col_a_header="", col_a_items=None,
                          col_b_header="", col_b_items=None,
                          key_message="", notes=""):
        """Pop-up: dark panel with 2-column comparison."""
        slide = self._new_slide()
        self._add_slide_title(slide, popup_title)
        self._add_layout_badge(slide, slide_num, "Pop-up on anchor")
        self._add_popup_dark_panel(slide)
        self._add_popup_header(slide, comp_num, comp_title)
        self._add_section_label(slide, "WHAT THIS COMPONENT COVERS",
                                POPUP_CONTENT_LEFT, POPUP_CONTENT_TOP,
                                color=COMPONENT_BLUE)
        col_w = Inches(5.5)
        ax = POPUP_CONTENT_LEFT
        self._add_section_label(slide, col_a_header, ax,
                                POPUP_CONTENT_TOP + Inches(0.5), color=LIGHT_GRAY)
        if col_a_items:
            self._add_bullets(slide, col_a_items,
                              ax + Inches(0.2), POPUP_CONTENT_TOP + Inches(0.9),
                              col_w, Inches(2.5), size=Pt(13),
                              bullet_color=COMPONENT_BLUE, text_color=WHITE)
        bx = Inches(7.0)
        self._add_section_label(slide, col_b_header, bx,
                                POPUP_CONTENT_TOP + Inches(0.5), color=LIGHT_GRAY)
        if col_b_items:
            self._add_bullets(slide, col_b_items,
                              bx + Inches(0.2), POPUP_CONTENT_TOP + Inches(0.9),
                              col_w, Inches(2.5), size=Pt(13),
                              bullet_color=COMPONENT_BLUE, text_color=WHITE)
        if key_message:
            self._add_section_label(slide, "KEY MESSAGE",
                                    POPUP_CONTENT_LEFT, Inches(5.5),
                                    color=COMPONENT_BLUE)
            self._add_text(slide, key_message, POPUP_CONTENT_LEFT, Inches(5.8),
                           POPUP_CONTENT_WIDTH, Inches(0.6),
                           size=Pt(14), color=LIGHT_GRAY)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    def add_popup_grid(self, slide_num, popup_title, comp_num, comp_title,
                       key_message="", grid_items=None, notes=""):
        """Pop-up: dark panel with 2x2 grid.
        grid_items: list of (header, body) tuples."""
        slide = self._new_slide()
        self._add_slide_title(slide, popup_title)
        self._add_layout_badge(slide, slide_num, "Pop-up on anchor")
        self._add_popup_dark_panel(slide)
        self._add_popup_header(slide, comp_num, comp_title)
        if key_message:
            self._add_section_label(slide, "KEY MESSAGE",
                                    POPUP_CONTENT_LEFT, POPUP_CONTENT_TOP,
                                    color=COMPONENT_BLUE)
            self._add_text(slide, key_message,
                           POPUP_CONTENT_LEFT, POPUP_CONTENT_TOP + Inches(0.3),
                           POPUP_CONTENT_WIDTH, Inches(0.5),
                           size=Pt(13), color=LIGHT_GRAY)
        if grid_items:
            grid_w = Inches(5.5)
            grid_h = Inches(1.3)
            gap_x = Inches(0.4)
            gap_y = Inches(0.3)
            start_y = POPUP_CONTENT_TOP + Inches(1.0)
            for i, (header, body) in enumerate(grid_items[:4]):
                col = i % 2
                row = i // 2
                x = POPUP_CONTENT_LEFT + col * (grid_w + gap_x)
                y = start_y + row * (grid_h + gap_y)
                self._add_content_box(slide, x, y, grid_w, grid_h,
                                      fill_color=PANEL_BG, border_color=BORDER_BLUE)
                self._add_section_label(slide, header, x + Inches(0.2),
                                        y + Inches(0.1), width=grid_w,
                                        color=COMPONENT_BLUE)
                self._add_text(slide, body, x + Inches(0.2), y + Inches(0.4),
                               grid_w - Inches(0.4), grid_h - Inches(0.5),
                               size=Pt(13), color=LIGHT_GRAY)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Legacy alias ─────────────────────────────────────────────

    def add_L12(self, slide_num, title, tiles, notes=""):
        return self.add_anchor(slide_num, title, tiles, notes=notes)

    # ── Public: Layout L13 — Objective / Agenda ──────────────────

    def add_L13(self, slide_num, title, objectives, notes=""):
        """L13 objectives — matches CSS: bg-secondary bar + blue left border + blue counter."""
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L13")
        bar_w = Inches(10.0)
        bar_h = Inches(0.65)
        gap = Inches(0.15)
        y = Inches(1.65)
        for i, obj in enumerate(objectives):
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         LEFT_MARGIN, y, bar_w, bar_h)
            bar.fill.solid()
            bar.fill.fore_color.rgb = BG_SECONDARY
            bar.line.fill.background()
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                            LEFT_MARGIN, y, Pt(4), bar_h)
            accent.fill.solid()
            accent.fill.fore_color.rgb = BRAND_BLUE
            accent.line.fill.background()
            self._add_text(slide, f"{i+1:02d}",
                           LEFT_MARGIN + Inches(0.2), y,
                           Inches(0.6), bar_h,
                           size=Pt(18), color=BRAND_BLUE, bold=True)
            self._add_text(slide, obj,
                           LEFT_MARGIN + Inches(0.8), y,
                           bar_w - Inches(1.0), bar_h,
                           size=Pt(15), color=WHITE)
            y += bar_h + gap
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Public: Layout L14 — Summary ─────────────────────────────

    def add_L14(self, slide_num, title, takeaways, cta="Continue", notes=""):
        """L14 summary — matches CSS: bg-secondary bar + yellow left border."""
        slide = self._new_slide()
        self._add_slide_title(slide, title)
        self._add_layout_badge(slide, slide_num, "L14")
        bar_w = Inches(10.0)
        bar_h = Inches(0.7)
        gap = Inches(0.15)
        y = Inches(1.65)
        for item in takeaways:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         LEFT_MARGIN, y, bar_w, bar_h)
            bar.fill.solid()
            bar.fill.fore_color.rgb = BG_SECONDARY
            bar.line.fill.background()
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                            LEFT_MARGIN, y, Pt(4), bar_h)
            accent.fill.solid()
            accent.fill.fore_color.rgb = YELLOW_ACCENT
            accent.line.fill.background()
            self._add_text(slide, item,
                           LEFT_MARGIN + Inches(0.3), y,
                           bar_w - Inches(0.5), bar_h,
                           size=Pt(15), color=WHITE)
            y += bar_h + gap
        btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     Inches(10.5), Inches(6.2), Inches(2), Inches(0.5))
        btn.fill.solid()
        btn.fill.fore_color.rgb = BRAND_BLUE
        btn.line.fill.background()
        self._add_text(slide, cta, Inches(10.5), Inches(6.2),
                       Inches(2), Inches(0.5), size=Pt(14),
                       color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        if notes:
            self._add_speaker_notes(slide, notes)
        self._add_footer(slide)
        return slide

    # ── Save ─────────────────────────────────────────────────────

    def save(self, path):
        self.prs.save(path)
        print(f"Storyboard saved: {path} ({self.slide_count} slides)")
